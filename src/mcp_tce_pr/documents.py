"""Extração limitada de documentos públicos; nunca executa conteúdo recebido."""

import asyncio
import csv
import io
import json
import zipfile

from bs4 import BeautifulSoup

from .client import SourceError, decode
from .portal import provenance, request


def extract(body: bytes, kind: str, page: int, limit: int) -> dict:
    if body.startswith(b"%PDF") or "pdf" in kind:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(body))
        if reader.is_encrypted:
            raise SourceError("PDF protegido; não será descriptografado.")
        count = len(reader.pages)
        text = "\n\n".join(
            f"[Página {i + 1}]\n{reader.pages[i].extract_text() or ''}"
            for i in range(page - 1, min(page - 1 + limit, count))
        )
        return {
            "formato": "pdf",
            "texto": text,
            "total_paginas": count,
            "proxima_pagina": page + limit if page + limit <= count else None,
            "aviso_pdf": "PDF digitalizado pode exigir OCR; texto vazio não significa documento vazio.",
        }
    if body.startswith(b"PK"):
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise SourceError("Documento descompactado excede 100 MiB.")
            names = archive.namelist()
            if "word/document.xml" in names:
                data = archive.read("word/document.xml")
                soup = BeautifulSoup(data, "xml")
                return {"formato": "docx", "texto": soup.get_text("\n", strip=True)}
            if "xl/workbook.xml" in names:
                from openpyxl import load_workbook

                wb = load_workbook(
                    io.BytesIO(body), read_only=True, data_only=True, keep_links=False
                )
                try:
                    sheets = []
                    for sheet in wb.worksheets[:20]:
                        rows = []
                        for row in sheet.iter_rows(
                            min_row=(page - 1) * 100 + 1,
                            max_row=page * 100,
                            max_col=50,
                            values_only=True,
                        ):
                            rows.append([str(v) if v is not None else "" for v in row])
                        sheets.append({"planilha": sheet.title, "linhas": rows})
                    return {
                        "formato": "xlsx",
                        "texto": json.dumps(sheets, ensure_ascii=False),
                        "planilhas": wb.sheetnames,
                        "faixa_linhas": [(page - 1) * 100 + 1, page * 100],
                        "aviso_planilha": "Até 20 planilhas, 50 colunas e 100 linhas por página; fórmulas usam valor salvo.",
                    }
                finally:
                    wb.close()
            raise SourceError("Arquivo ZIP: use listar_arquivos_zip_pr e consultar_csv_zip_pr.")
    if any(t in kind.lower() for t in ("text", "json", "xml", "csv")):
        text = decode(body)
        if "html" in kind:
            raise SourceError("URL retornou HTML em vez do documento. Use ler_pagina_portal_pr.")
        return {"formato": kind, "texto": text}
    # CSVs do TCE usam application/octet-stream.
    if b";" in body[:4096] and b"\x00" not in body[:4096]:
        csv.reader(io.StringIO(decode(body)), delimiter=";")
        return {"formato": "csv", "texto": decode(body)}
    raise SourceError(
        "Formato sem extrator: use o link original. P7S, XLS binário e imagens não são extraídos."
    )


async def ler_documento_pr(
    url: str, pagina: int = 1, paginas: int = 3, deslocamento: int = 0, limite: int = 20000
) -> dict:
    """Lê PDF, DOCX, XLSX, CSV, JSON, XML e texto de URLs públicas do TCE-PR.

    PDF: pagina inicial e até 5 paginas. XLSX: pagina seleciona blocos de 100 linhas.
    Texto extraído é paginado por deslocamento/limite; download máximo 20 MiB.
    Não faz OCR nem valida assinaturas P7S. Resposta inclui fonte e limites.
    """
    if pagina < 1 or not 1 <= paginas <= 5 or deslocamento < 0 or not 100 <= limite <= 50000:
        raise SourceError("Paginação inválida.")
    body, final, headers = await request(url, limit=20 * 1024 * 1024)
    result = await asyncio.to_thread(
        extract, body, headers.get("content-type", ""), pagina, paginas
    )
    text = result.pop("texto")
    return {
        **result,
        "texto": text[deslocamento : deslocamento + limite],
        "total_caracteres_extraidos": len(text),
        "proximo_deslocamento": deslocamento + limite
        if deslocamento + limite < len(text)
        else None,
        **provenance(final, headers),
    }
